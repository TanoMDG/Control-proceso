from openpyxl import Workbook
from io import BytesIO

from app.services.importer import REQUIRED_SOURCE_SHEETS, validate_excel_source
from tests.conftest import auth


def test_import_preview_requires_all_source_sheets(tmp_path):
    path = tmp_path / "incompleto.xlsx"
    workbook = Workbook()
    workbook.active.title = "17_Limites"
    workbook.active.append(["encabezado"])
    workbook.save(path)
    preview = validate_excel_source(path)
    assert preview["valid"] is False
    assert "18_Listas" in preview["missing_sheets"]


def test_import_preview_accepts_complete_nonempty_source_shape(tmp_path):
    path = tmp_path / "TEST_fuente.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in REQUIRED_SOURCE_SHEETS:
        sheet = workbook.create_sheet(name)
        sheet.append(["encabezado"])
        sheet.append(["TEST"])
    workbook.save(path)
    preview = validate_excel_source(path)
    assert preview["valid"] is True
    assert preview["sha256"]


def test_import_preview_accepts_source_sheet_with_extra_spacing(tmp_path):
    path = tmp_path / "TEST_fuente.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in REQUIRED_SOURCE_SHEETS:
        sheet = workbook.create_sheet("Tabla  conversion altura-tn" if name == "Tabla conversion altura-tn" else name)
        sheet.append(["encabezado"])
        sheet.append(["TEST"])
    workbook.save(path)
    assert validate_excel_source(path)["valid"] is True


def test_admin_import_preview_is_audit_logged(client, admin_token):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in REQUIRED_SOURCE_SHEETS:
        sheet = workbook.create_sheet(name)
        sheet.append(["encabezado"])
        sheet.append(["TEST"])
    content = BytesIO()
    workbook.save(content)
    response = client.post("/api/v1/importaciones/preview", headers=auth(admin_token), files={"file": ("TEST_fuente.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200
    assert response.json()["valid"] is True
    assert any(row["tabla"] == "importacion_datos" for row in client.get("/api/v1/auditoria", headers=auth(admin_token)).json())
