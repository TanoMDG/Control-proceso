from __future__ import annotations

import hashlib
import json
from pathlib import Path
from datetime import date
from decimal import Decimal
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.core import Catalog, ImportResult, ImportRun, Limit, LimitVersion, Person, PersonPosition, ReactionPlan, Role, SiloScale, SiloScalePoint, User
from app.services.audit import write_audit
from app.services.bootstrap import ensure_base_roles
from app.services.security import hash_password

REQUIRED_SOURCE_SHEETS = {
    "17_Limites",
    "18_Listas",
    "19_Personas",
    "10_Cod_Paradas",
    "15_Plan_Reaccion",
    "Tabla conversion altura-tn",
}


def canonical_sheet_name(name: str) -> str:
    return " ".join(name.split())


def validate_excel_source(path: Path) -> dict:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return {"valid": False, "error": "El archivo debe ser Excel .xlsx o .xlsm"}
    content = path.read_bytes()
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets_by_canonical_name = {canonical_sheet_name(name): name for name in workbook.sheetnames}
    missing = sorted(REQUIRED_SOURCE_SHEETS - set(sheets_by_canonical_name))
    rows_by_sheet = {canonical: workbook[sheets_by_canonical_name[canonical]].max_row for canonical in sorted(REQUIRED_SOURCE_SHEETS & set(sheets_by_canonical_name))}
    errors = [f"La hoja requerida no existe: {sheet}" for sheet in missing]
    errors.extend(f"La hoja requerida esta vacia: {sheet}" for sheet, rows in rows_by_sheet.items() if rows < 2)
    return {
        "filename": path.name,
        "sha256": hashlib.sha256(content).hexdigest(),
        "valid": not errors,
        "sheets_found": sorted(workbook.sheetnames),
        "missing_sheets": missing,
        "rows_by_sheet": rows_by_sheet,
        "errors": errors,
    }


def _text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _decimal(value) -> Decimal | None:
    return Decimal(str(value)) if value is not None else None


def _source_sheet(workbook, canonical_name: str):
    return workbook[{canonical_sheet_name(name): name for name in workbook.sheetnames}[canonical_name]]


def _upsert_catalog(db: Session, *, tipo: str, codigo: str, descripcion: str, atributos: dict | None = None) -> Catalog:
    item = db.scalar(select(Catalog).where(Catalog.tipo == tipo, Catalog.codigo == codigo))
    if item is None:
        item = Catalog(tipo=tipo, codigo=codigo, descripcion=descripcion, atributos=atributos, activo=True, fecha_baja=None)
        db.add(item)
    else:
        item.descripcion, item.atributos, item.activo, item.fecha_baja = descripcion, atributos, True, None
    db.flush()
    return item


def ensure_development_import_admin(db: Session, *, password: str) -> User:
    existing = db.scalar(select(User).where(User.nombre_usuario == "admin.dev"))
    if existing is not None:
        return existing
    roles = ensure_base_roles(db)
    person = db.scalar(select(Person).where(Person.legajo == "99999"))
    if person is None:
        person = Person(legajo="99999", apellido_nombre="Administrador Desarrollo", activo=True, fecha_baja=None, es_cuenta_tecnica_dev=True)
        db.add(person)
        db.flush()
    elif not person.es_cuenta_tecnica_dev:
        raise ValueError("El legajo 99999 ya pertenece a una persona no técnica")
    user = User(id_persona=person.id, nombre_usuario="admin.dev", password_hash=hash_password(password), id_rol=roles["ADMIN"].id, sector="Administracion", activo=True)
    db.add(user)
    db.flush()
    write_audit(db, user_id=user.id, table="usuario", record_id=str(user.id), action="ALTA_TECNICA_DEV", after={"usuario": "admin.dev", "legajo": "99999", "entorno": "development"})
    return user


def import_master_data(db: Session, *, path: Path, actor: User, vigente_desde: date | None = None) -> dict:
    preview = validate_excel_source(path)
    if not preview["valid"]:
        raise ValueError("El libro fuente no supera el dry-run")
    effective_date = vigente_desde or date.today()
    workbook = load_workbook(path, read_only=True, data_only=True)
    run = ImportRun(nombre_archivo=path.name, sha256=preview["sha256"], modo="APLICAR", estado="EN_PROCESO", id_usuario=actor.id, resumen={})
    db.add(run)
    db.flush()
    counts = {"limites": 0, "catalogos": 0, "personas": 0, "puestos": 0, "desvios": 0, "escalas": 0, "puntos_escala": 0, "discrepancias": 0}

    listas = _source_sheet(workbook, "18_Listas")
    values = list(listas.iter_rows(min_row=4, values_only=True))
    for row in values:
        turno, box, linea, prensa, linea_prensa, formato, cavidad, sector, equipo, silo, _, limpieza, _, _, formato_cm, espesor, tolerancia = row[:17]
        for tipo, value in (("turno", turno), ("box", box), ("linea", linea), ("cavidad", cavidad), ("sector", sector), ("equipo", equipo), ("silo", silo), ("tipo_limpieza", limpieza)):
            if value is not None:
                _upsert_catalog(db, tipo=tipo, codigo=_text(value), descripcion=_text(value))
                counts["catalogos"] += 1
        if prensa is not None:
            _upsert_catalog(db, tipo="prensa", codigo=_text(prensa), descripcion=_text(prensa), atributos={"linea": _text(linea_prensa)})
            counts["catalogos"] += 1
        if formato is not None:
            _upsert_catalog(db, tipo="formato", codigo=_text(formato), descripcion=_text(formato), atributos={"formato_cm": _text(formato_cm), "espesor_nominal_mm": _text(espesor), "tolerancia_mm": _text(tolerancia)})
            counts["catalogos"] += 1

    paradas = _source_sheet(workbook, "10_Cod_Paradas")
    for code, cause, group, _ in paradas.iter_rows(min_row=4, max_col=4, values_only=True):
        if code and str(code).startswith("P"):
            _upsert_catalog(db, tipo="cod_parada", codigo=_text(code), descripcion=_text(cause), atributos={"grupo": _text(group)})
            counts["catalogos"] += 1

    lineas = {item.codigo: item.id for item in db.scalars(select(Catalog).where(Catalog.tipo == "linea"))}
    personas = _source_sheet(workbook, "19_Personas")
    role_columns = {4: "operador_molienda", 5: "palero", 6: "prensero", 7: "supervisor", 8: "laboratorio", 9: "mantenimiento", 10: "administracion"}
    for row in personas.iter_rows(min_row=4, max_col=15, values_only=True):
        if not row[0]:
            continue
        legajo, nombre, abreviatura, linea_habitual = (_text(row[0]), _text(row[1]), _text(row[2]), _text(row[3]))
        person = db.scalar(select(Person).where(Person.legajo == legajo))
        if person is None:
            person = Person(legajo=legajo, apellido_nombre=nombre, abreviatura=abreviatura, id_linea_habitual=lineas.get(linea_habitual), activo=_text(row[12]) == "Sí", fecha_baja=row[13])
            db.add(person)
            db.flush()
            counts["personas"] += 1
        for index, puesto in role_columns.items():
            if _text(row[index]) == "X":
                existing = db.scalar(select(PersonPosition).where(PersonPosition.id_persona == person.id, PersonPosition.puesto == puesto, PersonPosition.vigente_hasta.is_(None)))
                if existing is None:
                    db.add(PersonPosition(id_persona=person.id, puesto=puesto, id_linea=lineas.get(linea_habitual), vigente_desde=effective_date, vigente_hasta=None))
                    counts["puestos"] += 1

    limites = _source_sheet(workbook, "17_Limites")
    for row in limites.iter_rows(min_row=4, max_col=11, values_only=True):
        if not row[0] or not str(row[0]).startswith("L"):
            continue
        limit_id = _text(row[0])
        limit = db.get(Limit, limit_id)
        if limit is None:
            limit = Limit(id=limit_id, variable=_text(row[1]), etapa=_text(row[2]), unidad=_text(row[3]), tipo_dato=_text(row[7]), referencia_fuente=_text(row[6]), nota_fuente=_text(row[10]))
            db.add(limit)
        else:
            limit.variable, limit.etapa, limit.unidad, limit.tipo_dato = _text(row[1]), _text(row[2]), _text(row[3]), _text(row[7])
            limit.referencia_fuente, limit.nota_fuente = _text(row[6]), _text(row[10])
        db.flush()
        version = db.scalar(select(LimitVersion).where(LimitVersion.id_limite == limit_id, LimitVersion.vigente_desde == effective_date))
        if version is None:
            # F0 preserves the source values; operational warning/deviation mapping is configured in its formal module phase.
            db.add(LimitVersion(id_limite=limit_id, vigente_desde=effective_date, valor_min=_decimal(row[4]), valor_max=_decimal(row[5]), nivel="INFORMATIVO", motivo_cambio="Importación inicial desde libro fuente", id_usuario_alta=actor.id, estado="VIGENTE"))
            counts["limites"] += 1

    overrides_path = Path(__file__).parent.parent / "config" / "v1_4_final_reaction_overrides.json"
    overrides = json.loads(overrides_path.read_text(encoding="utf-8"))
    reactions = _source_sheet(workbook, "15_Plan_Reaccion")
    for row in reactions.iter_rows(min_row=4, max_col=11, values_only=True):
        if not row[0] or not str(row[0]).startswith("D"):
            continue
        deviation_id = _text(row[0])
        source_escalation = _text(row[8])
        override = overrides.get(deviation_id)
        effective_escalation = override["registro_escalamiento"] if override else source_escalation
        plan = db.get(ReactionPlan, deviation_id)
        values = {"senal": _text(row[1]), "etapa": _text(row[2]), "limite_referencia": _text(row[3]), "causas_probables": _text(row[4]), "accion_inmediata": _text(row[5]), "verificacion": _text(row[6]), "plazo_texto": _text(row[7]), "registro_escalamiento": effective_escalation, "plazo_horas": _decimal(row[9]), "tipo_plazo": _text(row[10]), "texto_excel_original": source_escalation if override else None, "motivo_prevalencia": override["motivo"] if override else None}
        if plan is None:
            db.add(ReactionPlan(id_desvio=deviation_id, **values))
            counts["desvios"] += 1
        else:
            for field, value in values.items():
                setattr(plan, field, value)
        if override:
            db.add(ImportResult(id_importacion=run.id, hoja="15_Plan_Reaccion", fila=int(deviation_id[1:]) + 3, nivel="DISCREPANCIA", mensaje=override["motivo"], valor_original=source_escalation, valor_normativo=effective_escalation))
            counts["discrepancias"] += 1

    scales = _source_sheet(workbook, "Tabla conversion altura-tn")
    scale_values = list(scales.iter_rows(min_row=3, max_row=24, max_col=3, values_only=True))
    for column, group in ((1, "1-8"), (2, "9-16")):
        scale = db.scalar(select(SiloScale).where(SiloScale.grupo_silos == group, SiloScale.version == 1))
        if scale is None:
            scale = SiloScale(grupo_silos=group, version=1, vigente_desde=effective_date, aprobada_por=actor.id)
            db.add(scale)
            db.flush()
            counts["escalas"] += 1
        for height, first, second in scale_values:
            tonnes = first if column == 1 else second
            point = db.scalar(select(SiloScalePoint).where(SiloScalePoint.id_escala == scale.id, SiloScalePoint.altura_m == _decimal(height)))
            if point is None:
                db.add(SiloScalePoint(id_escala=scale.id, altura_m=_decimal(height), toneladas=_decimal(tonnes)))
                counts["puntos_escala"] += 1

    run.estado = "APLICADO"
    run.resumen = {**preview, "counts": counts, "vigente_desde": effective_date.isoformat()}
    write_audit(db, user_id=actor.id, table="importacion_datos", record_id=str(run.id), action="IMPORTACION_APLICADA", after=run.resumen, reason="Carga inicial de datos maestros desde Excel")
    return {"run_id": str(run.id), **run.resumen}
